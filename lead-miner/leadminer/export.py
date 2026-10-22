"""Write leads to CSV (opens in Excel with Arabic intact) and XLSX."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import Lead


def _row(lead: Lead, fields: list[str]) -> list:
    return [getattr(lead, f) if getattr(lead, f) is not None else "" for f in fields]


def write_csv(path: Path, leads: list[Lead], fields: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig adds the BOM Excel needs to show Arabic correctly.
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(fields)
        writer.writerows(_row(lead, fields) for lead in leads)
    return path


def write_xlsx(path: Path, leads: list[Lead], fields: list[str]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "leads"
    ws.append(fields)
    for lead in leads:
        ws.append(_row(lead, fields))
    ws.freeze_panes = "A2"
    for i, field in enumerate(fields, start=1):
        width = max([len(field)] + [len(str(c.value or "")) for c in ws[get_column_letter(i)][1:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    wb.save(path)
    return path
